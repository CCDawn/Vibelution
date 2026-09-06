r"""挑战杯高质量假说包 → 脱敏、可自校验的逐题提交结果。

输入：SCI-001.md 至 SCI-125.md，以及只读完成账本和官方题库。
输出：summaries、projections 和 manifest。

原则：公共结果固定投影假说生成与待执行研究计划，不投影内部实验记录或附件；同时不输出
Agent、会话、机器路径或原始内部附录。解析与语义门失败时停止，不得以不完整结果覆盖
正式交付目录。
"""

import argparse
import hashlib
import json
import os
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CHALLENGE_ROOT = Path(
    os.environ.get("VIBELUTION_CHALLENGE_CUP_ROOT", Path.home() / "Desktop" / "挑战杯")
)
TZ = timezone(timedelta(hours=8))
_CANDIDATE_TOKEN = r"[A-Z](?:[′’'])?"


def normalize_candidate_key(value):
    return str(value or "").upper().replace("'", "′").replace("’", "′")


def extract_candidate_key(value):
    text = str(value or "").strip()
    match = re.search(
        rf"候选\s*\**\s*({_CANDIDATE_TOKEN})\s*\**", text, re.IGNORECASE
    )
    if not match:
        match = re.match(
            rf"\**\s*({_CANDIDATE_TOKEN})\s*\**(?=$|[\s，,：:。.)（(])",
            text,
            re.IGNORECASE,
        )
    return normalize_candidate_key(match.group(1)) if match else ""


def extract_header_candidate_key(value):
    key = extract_candidate_key(value)
    if key:
        return key
    match = re.match(
        rf"\**\s*({_CANDIDATE_TOKEN})(?=\s|\**$)",
        str(value or "").strip(),
        re.IGNORECASE,
    )
    return normalize_candidate_key(match.group(1)) if match else ""


def candidate_base_key(value):
    key = normalize_candidate_key(value)
    return key[:1]


def resolve_candidate_key(reference, candidates):
    reference = normalize_candidate_key(reference)
    candidate_keys = [candidate["letter"] for candidate in candidates]
    if reference in candidate_keys:
        return reference
    base = candidate_base_key(reference)
    base_matches = [key for key in candidate_keys if candidate_base_key(key) == base]
    return base_matches[0] if len(base_matches) == 1 else ""


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def now_iso():
    return datetime.now(TZ).strftime("%Y-%m-%dT%H:%M:%S+08:00")


def clean_question(raw):
    q = re.split(r"（来源|\(来源", raw)[0]
    return q.replace("`", "").strip()


def safe_clip(value, limit):
    """在句子/分句边界裁剪，避免裸字符切片产生句中截断。"""
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if len(text) <= limit:
        return text
    window = text[: limit + 1]
    floor = max(0, int(limit * 0.55))
    cuts = [
        m.end() for m in re.finditer(r"[。！？；.!?;]\s*", window) if m.end() >= floor
    ]
    if not cuts:
        cuts = [m.start() for m in re.finditer(r"\s+", window) if m.start() >= floor]
    end = cuts[-1] if cuts else limit
    return window[:end].rstrip(" ，、:：") + "…"


def esc_cell(value, limit=None):
    """表格单元格专用：裁剪后转义裸竖线，防止 |ρ| 类数学记号撑破 Markdown 表格列位。"""
    text = (
        safe_clip(value, limit)
        if limit is not None
        else re.sub(r"\s+", " ", str(value or "")).strip()
    )
    return text.replace("|", "\\|")


def normalize_source_url(locator):
    """将显式 URL 或 DOI 定位符统一为可点击 URL。"""
    raw = str(locator or "").strip()
    url = re.search(r"https?://[^\s，。；（]+", raw, re.IGNORECASE)
    if url:
        normalized = _trim_locator(url.group(0))
        return None if _is_placeholder_locator(normalized) else normalized
    doi = re.search(
        r"(?:doi\s*[:：]?\s*)?(10\.\d{4,9}/[^\s\"'，。；（]+)", raw, re.IGNORECASE
    )
    if doi:
        normalized = "https://doi.org/" + _trim_locator(doi.group(1))
        return None if _is_placeholder_locator(normalized) else normalized
    pmc = re.search(r"\b(PMC\d+)\b", raw, re.IGNORECASE)
    if pmc:
        return f"https://pmc.ncbi.nlm.nih.gov/articles/{pmc.group(1).upper()}/"
    pmid = re.search(r"\bPMID\s*[:：]?\s*(\d{6,9})\b", raw, re.IGNORECASE)
    if pmid:
        return f"https://pubmed.ncbi.nlm.nih.gov/{pmid.group(1)}/"
    arxiv = re.search(
        r"\barXiv\s*[:：]?\s*((?:\d{4}\.\d{4,5})|(?:[a-z-]+/\d{7}))(?:v\d+)?\b",
        raw,
        re.IGNORECASE,
    )
    if arxiv:
        return f"https://arxiv.org/abs/{arxiv.group(1)}"
    bookshelf = re.search(r"\b(NBK\d+)\b", raw, re.IGNORECASE)
    if bookshelf:
        return f"https://www.ncbi.nlm.nih.gov/books/{bookshelf.group(1).upper()}/"
    domain = re.search(
        r"(?<!@)\b(?:www\.)?[A-Z0-9.-]+\.[A-Z]{2,}(?:/[^\s，。；（]*)?",
        raw,
        re.IGNORECASE,
    )
    if domain:
        normalized = "https://" + _trim_locator(domain.group(0))
        return None if _is_placeholder_locator(normalized) else normalized
    return None


def _trim_locator(value):
    locator = str(value or "").rstrip(".,，。；）】")
    while locator.endswith(")") and locator.count(")") > locator.count("("):
        locator = locator[:-1]
    while locator.endswith("]") and locator.count("]") > locator.count("["):
        locator = locator[:-1]
    return locator


def _is_placeholder_locator(value):
    locator = str(value or "")
    return "…" in locator or re.search(r"(?:^|/)\.\.\.(?:/|$)", locator) is not None


def load_catalog_questions(catalog_path):
    mapping = {}
    with open(catalog_path, encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, list):
        items = data
    elif isinstance(data, dict):
        items = data.get("questions") or data.get("items") or []
    else:
        raise TypeError("catalog root must be a list or object")
    for q in items:
        if not isinstance(q, dict):
            continue
        qid = q.get("id") or q.get("question_id") or q.get("qid")
        text = q.get("question") or q.get("question_en") or q.get("text") or ""
        if qid and text:
            mapping[str(qid)] = str(text).strip()
    return mapping


def ledger_done_questions(ledger_path):
    from openpyxl import load_workbook

    wb = load_workbook(ledger_path, read_only=True)
    ws = wb["125题进度账本"]
    done = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] and row[5] in ("☑", "\u2611"):
            done.add(str(row[0]))
    wb.close()
    return done


def split_sections(text):
    chinese = {
        "一": 1,
        "二": 2,
        "三": 3,
        "四": 4,
        "五": 5,
        "六": 6,
        "七": 7,
        "八": 8,
        "九": 9,
        "十": 10,
    }
    marks = list(
        re.finditer(
            r"^##\s+(?:(\d+)\s*[.、]|([一二三四五六七八九十])、)\s*.*$",
            text,
            re.MULTILINE,
        )
    )
    secs = {}
    for i, m in enumerate(marks):
        start = m.end()
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        number = int(m.group(1)) if m.group(1) else chinese[m.group(2)]
        secs[number] = text[start:end].strip()
    return secs


def parse_meta(text):
    meta = {}
    aliases = {
        "题目原文": ["题目原文", "题目（EN）"],
        "领域": ["领域"],
        "完成时间": ["完成时间", "执行时间"],
        "执行 Agent": ["执行 Agent"],
    }
    for key, labels in aliases.items():
        for label in labels:
            m = re.search(
                rf"^-\s*(?:\*\*)?{re.escape(label)}(?:\*\*)?\s*[：:]\s*(.+)$",
                text,
                re.MULTILINE,
            )
            if m:
                meta[key] = m.group(1).strip()
                break
    return meta


def slice_fields(body, labels):
    """按字段标签切片正文。标签可带（括注）。重叠/嵌套标签保留最长匹配。"""
    matches = []
    for lab in labels:
        prediction_suffix = (
            r"(?:\s*P\d+(?:（[^）]*）)?)?"
            if lab in ("可检验预测", "预测")
            else ""
        )
        for m in re.finditer(
            rf"(?:^|\n|[；;]\s*)(?:[-*]\s*)?\**({re.escape(lab)}){prediction_suffix}\**\s*(（[^）]*）)?\s*\**\s*[：:=＝]",
            body,
        ):
            actual = m.group(1) + (m.group(2) or "")
            matches.append((m.start(1), m.end(), actual, lab))
    matches.sort(key=lambda x: (x[0], -x[1]))
    starts = []
    last_end = -1
    for s, e, actual, base in matches:
        if s >= last_end:
            starts.append((s, e, actual, base))
            last_end = e
    fields = {}
    for i, (s, e, actual, base) in enumerate(starts):
        end_pos = starts[i + 1][0] if i + 1 < len(starts) else len(body)
        val = body[e:end_pos].strip().strip("*").strip()
        val = re.sub(r"\s*\n\s*", " ", val)
        val = val.rstrip(" \t\r\n-*；;")
        if val:
            fields.setdefault(actual, []).append(val)
            if base != actual:
                fields.setdefault(base, []).append(val)
    return fields


def extract_predictions(body):
    """提取『可检验预测 P1（…）：值』式逐行预测。"""
    preds = []
    lines = body.splitlines()
    index = 0
    while index < len(lines):
        ln = lines[index]
        s = ln.strip()
        m = re.match(
            r"(?:[-*]\s*)?\**(?:可检验预测|预测)\s*(?:P\d(?:（[^）]*）)?)?\**\s*[：:]\s*(.*)$",
            s,
        )
        if m:
            v = m.group(1).strip().strip("*").strip()
            if v:
                preds.append(v)
                index += 1
                continue
            following_index = index + 1
            numbered = []
            while following_index < len(lines):
                following = lines[following_index].strip()
                if not following:
                    following_index += 1
                    continue
                if re.match(r"(?:#{1,6}\s+|\**(?:证伪|适用边界|不确定性))", following):
                    break
                numbered_match = re.match(r"\d+[.、)]\s*(.+)$", following)
                if numbered_match:
                    numbered.append(numbered_match.group(1).strip())
                    following_index += 1
                    continue
                if not numbered:
                    numbered.append(following.lstrip(">-").strip())
                break
            preds.extend(value for value in numbered if value)
            index = following_index
            continue
        m = re.match(r"(?:[-*]\s*)?\**P\d+[^：:]*\**\s*[：:]\s*(.+)$", s, re.IGNORECASE)
        if m:
            v = m.group(1).strip().strip("*").strip()
            if v:
                preds.append(v)
        index += 1
    return list(dict.fromkeys(preds))


def parse_sources(sec):
    rows = []
    for line in sec.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip().strip("*") for c in line.strip("|").split("|")]
        if not cells or not re.match(r"^S\d+$", cells[0]):
            continue
        if len(cells) >= 7:
            locator = cells[5]
            rows.append(
                {
                    "sid": cells[0],
                    "title": cells[1],
                    "authors": cells[2],
                    "year": cells[3],
                    "stype": cells[4],
                    "url": locator,
                    "source_url": normalize_source_url(locator),
                    "contribution": cells[6],
                }
            )
        elif len(cells) >= 5:
            citation = cells[1]
            year_match = re.search(r"\b(19|20)\d{2}\b", citation)
            year = year_match.group(0) if year_match else ""
            doi = re.search(
                r"(?:doi\s*[:：]?\s*)?(10\.\d{4,9}/[-._;()/:A-Z0-9]+)",
                citation,
                re.IGNORECASE,
            )
            locator = doi.group(1).rstrip(".,，。；;") if doi else ""
            authors = citation.split("(", 1)[0].strip().rstrip(".")
            rows.append(
                {
                    "sid": cells[0],
                    "title": citation,
                    "authors": authors,
                    "year": year,
                    "stype": cells[2],
                    "url": locator,
                    "source_url": normalize_source_url(locator),
                    "contribution": cells[4],
                }
            )
    return rows


def parse_candidates(sec):
    heading_re = re.compile(
        rf"^(?:###\s*|\*\*)候选\s*\**\s*({_CANDIDATE_TOKEN})\s*\**\s*[：:·\-—–\s]*(.*?)(?:\*\*)?\s*$",
        re.MULTILINE,
    )
    marks = list(heading_re.finditer(sec))
    cands = []
    for i, m in enumerate(marks):
        letter = normalize_candidate_key(m.group(1))
        name = m.group(2).strip().strip("：:·-— ").strip()
        body = sec[m.end() : marks[i + 1].start() if i + 1 < len(marks) else len(sec)]
        f = slice_fields(body, ["陈述", "机制", "证据指向", "风险", "风险/薄弱点"])
        statement = (f.get("陈述") or [""])[0]
        if not statement:
            paras = [
                re.sub(r"\*+", "", x).strip()
                for x in re.split(r"\n\s*\n", body)
                if x.strip() and not x.lstrip().startswith(("-", "|", "#"))
            ]
            statement = paras[0] if paras else ""
        mechanism = (f.get("机制（因果链）") or f.get("机制") or [""])[0]
        if not mechanism:
            mm = re.search(r"(?:^|\n)\s*[-*]\s*\**机制\**\s*[：:]\s*(.+)", body)
            mechanism = mm.group(1).strip() if mm else ""
        cands.append(
            {
                "letter": letter,
                "name": name or "候选 " + letter,
                "statement": statement,
                "mechanism": mechanism,
                "evidence": (f.get("证据指向（S 编号）") or f.get("证据指向") or [""])[
                    0
                ],
                "risk": (f.get("风险/薄弱点") or f.get("风险") or [""])[0],
                "predictions": extract_predictions(body),
                "raw": safe_clip(body, 600),
            }
        )
    return cands


def parse_review(sec):
    blocks = []
    current = []
    for line in sec.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            if current:
                blocks.append(current)
                current = []
            continue
        cells = [c.strip().strip("*") for c in line.strip("|").split("|")]
        if cells and cells[0]:
            current.append(cells)
    if current:
        blocks.append(current)

    best = (None, [], None)
    best_rank = (-1, -1)
    for block in blocks:
        header = next(
            (
                cells
                for cells in block
                if cells[0] in ("维度", "评价维度", "评分维度")
            ),
            None,
        )
        total = next(
            (
                cells
                for cells in block
                if cells[0].startswith("总分") or cells[0] in ("合计", "总计")
            ),
            None,
        )
        if not header:
            continue
        candidate_count = sum(
            bool(extract_header_candidate_key(c)) for c in header[1:]
        )
        value_count = (
            sum(
                bool(re.fullmatch(r"\s*\d+(?:\.\d+)?\s*", c))
                for c in total[1:]
            )
            if total
            else 0
        )
        rank = (candidate_count, value_count)
        if rank <= best_rank:
            continue
        rows = [
            cells
            for cells in block
            if cells is not header
            and cells is not total
            and not set(cells[0]) <= set("-: ")
        ]
        best = (header, rows, total)
        best_rank = rank
    return best


def _heading_trigger(sec, pos):
    back = sec[max(0, pos - 400) : pos]
    lines = [l for l in back.splitlines() if l.strip()]
    for l in reversed(lines):
        m = re.search(r"[（(]([^（）()]{2,60})[）)]\s*$", l.strip())
        if m and ("修订" in l or "触发" in m.group(1)):
            return m.group(1)
    return ""


def parse_revisions(sec):
    headings = [
        match.start()
        for match in re.finditer(
            r"^#{2,6}\s+修订(?:\s*\d+|[一二三四五六七八九十])", sec, re.MULTILINE
        )
    ]
    triggers = [
        m.start()
        for m in re.finditer(
            r"(?:^|\n)\s*(?:[-*]\s*)?\**(?:修订触发|触发(?:原因|证据)?)(?:（[^）]*）)?\**\s*[：:]",
            sec,
        )
    ]
    before_markers = [
        m.start()
        for m in re.finditer(
            r"(?:^|\n)(?:[-*]\s*)?\**修订前(?:表述)?(?:（[^）]*）)?\**\s*[：:]", sec
        )
    ]
    if headings:
        trig = headings
    elif before_markers and (not triggers or before_markers[0] < triggers[0]):
        trig = before_markers
    else:
        trig = triggers
    if not trig:
        return []
    bounds = trig + [len(sec)]
    recs = []
    for i in range(len(trig)):
        chunk = sec[bounds[i] : bounds[i + 1]]
        f = slice_fields(
            chunk,
            [
                "修订触发",
                "触发原因",
                "触发证据",
                "触发",
                "修订前表述",
                "修订前",
                "修订内容",
                "修订后主假说",
                "修订后表述",
                "修订后",
                "为什么更好",
                "改进原因",
                "改进",
                "实质影响",
            ],
        )
        recs.append(
            {
                "trigger": (
                    f.get("修订触发")
                    or f.get("触发原因")
                    or f.get("触发证据")
                    or f.get("触发")
                    or [""]
                )[0]
                or _heading_trigger(sec, bounds[i]),
                "before": (f.get("修订前") or f.get("修订前表述") or [""])[0],
                "after": (
                    f.get("修订后主假说")
                    or f.get("修订后")
                    or f.get("修订后表述")
                    or f.get("修订内容")
                    or [""]
                )[0],
                "why": (
                    f.get("为什么更好")
                    or f.get("改进原因")
                    or f.get("改进")
                    or f.get("实质影响")
                    or [""]
                )[0],
            }
        )
    return recs


def parse_selection(sec):
    return slice_fields(
        sec,
        [
            "主假说",
            "主假设",
            "备选假说",
            "备份假说",
            "备选假设",
            "保留备选",
            "胜出原因",
            "落选原因",
        ],
    )


def parse_plan(sec):
    rows = []
    for line in sec.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip().strip("*") for c in line.strip("|").split("|")]
        if not cells or cells[0] in ("步骤", "关键步骤") or set(cells[0]) <= set("-: "):
            continue
        if re.match(r"^\d+", cells[0]):
            rows.append(cells)
    return rows


def parse_scores(header, total, review_sec, *, warnings=None):
    """按候选字母映射总分；正文显式总分只补缺，不覆盖结构化总分行。"""
    warnings = warnings if warnings is not None else []
    scores = {}
    candidate_columns = {}
    if header and total:
        for i, cname in enumerate(header[1:], 1):
            if re.search(r"理由|说明|结论|备注", cname):
                continue
            key = extract_header_candidate_key(cname)
            if key:
                candidate_columns[i] = key
            if key and i < len(total):
                value = re.search(r"\d+(?:\.\d+)?", total[i])
                if value:
                    scores[key] = value.group(0)

    _, rows, _ = parse_review(review_sec)
    sums = {key: 0.0 for key in candidate_columns.values()}
    counts = {key: 0 for key in candidate_columns.values()}
    for row in rows:
        for index, key in candidate_columns.items():
            if index >= len(row):
                continue
            value = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*(?:/\s*\d+(?:\.\d+)?)?\s*", row[index])
            if value:
                sums[key] += float(value.group(1))
                counts[key] += 1
    for key in candidate_columns.values():
        if counts[key] != 7:
            continue
        computed = sums[key]
        computed_text = str(int(computed)) if computed.is_integer() else str(computed)
        declared = scores.get(key)
        if declared is not None and float(declared) != computed:
            warnings.append(
                f"候选 {key} 七维明细求和 {computed_text} 与声明总分 {declared} 不一致（采用明细求和）"
            )
        scores[key] = computed_text

    for m in re.finditer(
        r"(?<![A-Z0-9−-])([A-F](?:[′’'])?)\s*[=＝]\s*(\d+(?:\.\d+)?)"
        r"(?!\s*[+＋])(?:\s*(?:分|/\s*35)?)?(?=$|[；;，,\n|])",
        review_sec,
        re.IGNORECASE,
    ):
        key = normalize_candidate_key(m.group(1))
        if candidate_columns and key not in candidate_columns.values():
            continue
        scores.setdefault(key, m.group(2))
    return scores


def infer_source_relation(source_type, contribution):
    text = f"{source_type or ''} {contribution or ''}".lower()
    if re.search(r"反证|反驳|冲突|不支持|否证", text):
        if re.search(r"边界|局限|限制|外推|适用范围", text):
            return "boundary"
        return "challenges"
    if re.search(r"边界|局限|限制|外推|适用范围", text):
        return "boundary"
    if re.search(r"方法|测量|量表|数据集|基准|协议|平台", text):
        return "method"
    return "supports"


def infer_verification_status(source_url):
    # URL normalization proves only that a locator can be represented.  A real
    # verification status requires a durable provider or human-review receipt,
    # which this exporter does not receive.
    return "unverified"


def build_dataset_targets(plan_rows):
    targets = []
    for i, cells in enumerate(plan_rows, 1):
        item = {"target_id": f"T{i}", "stage": cells[0], "status": "proposed"}
        if len(cells) >= 4:
            item.update(
                action=cells[1],
                data_or_experiment=cells[2],
                decision_criterion=cells[3],
            )
        elif len(cells) == 3:
            item.update(action=cells[1], decision_criterion=cells[2])
        elif len(cells) == 2:
            item.update(action=cells[1])
        targets.append(item)
    return targets or [
        {
            "status": "proposed_target_not_structured",
            "note": "研究计划表未解析；请以原始假说包附录为准。",
        }
    ]


def revision_category(trigger):
    t = trigger or ""
    if re.search(r"证据|来源|反证|S\d|文献", t):
        return "科学证据"
    if re.search(r"评审|比较|打分|筛选", t):
        return "假设比较与筛选"
    if re.search(r"计划|预测|研究步骤", t):
        return "研究计划"
    if re.search(r"边界|适用|外推", t):
        return "不确定性与边界"
    return "候选假设"


def build_outputs(qid, src_path, catalog_q, src_text):
    meta = parse_meta(src_text)
    secs = split_sections(src_text)
    warnings = []

    question = clean_question(meta.get("题目原文", catalog_q or ""))
    domain = meta.get("领域", "")

    s1f = slice_fields(
        secs.get(1, ""),
        ["复述", "问题复述", "关键术语", "研究边界", "题意存疑点", "题意存疑"],
    )
    sources = parse_sources(secs.get(2, ""))
    cands = parse_candidates(secs.get(3, ""))
    header, _review_rows, total = parse_review(secs.get(4, ""))
    revisions = parse_revisions(secs.get(5, ""))
    substantive_revisions = [
        revision
        for revision in revisions
        if revision.get("trigger")
        and revision.get("before")
        and revision.get("after")
        and revision.get("why")
    ]
    sel = parse_selection(secs.get(6, ""))
    main_pkg = slice_fields(
        secs.get(7, ""),
        [
            "假说陈述",
            "可证伪假说",
            "竞争解释",
            "替代解释",
            "可检验预测",
            "预测",
            "证伪条件",
            "证伪/退出条件",
            "证伪/削弱条件",
            "适用边界",
            "不确定性声明",
            "不确定性",
        ],
    )
    sec10 = slice_fields(
        secs.get(10, "") or "", ["paper_title", "paper_abstract", "结果分类"]
    )
    predictions = extract_predictions(secs.get(7, ""))
    plan_rows = parse_plan(secs.get(8, ""))

    if not sources:
        warnings.append("来源表未解析（原文附录兜底）")
    if not cands:
        warnings.append("候选假说未解析（原文附录兜底）")
    if not revisions:
        warnings.append("修订记录未解析（原文附录兜底）")
        revisions = [{"trigger": "", "before": "", "after": "", "why": ""}]
    if not plan_rows:
        warnings.append("研究计划表未解析（原文附录兜底）")

    main_hyp = (sel.get("主假说") or sel.get("主假设") or [""])[0]
    backup_hyp = (
        sel.get("备选假说")
        or sel.get("备份假说")
        or sel.get("备选假设")
        or sel.get("保留备选")
        or [""]
    )[0]
    main_statement = (
        main_pkg.get("假说陈述（可证伪）")
        or main_pkg.get("假说陈述")
        or main_pkg.get("可证伪假说")
        or [""]
    )[0]
    competing_explanation = (
        main_pkg.get("竞争解释") or main_pkg.get("替代解释") or [""]
    )[0]
    selected_ref = extract_candidate_key(main_hyp)
    selected_letter = resolve_candidate_key(selected_ref, cands)
    if not selected_letter and cands:
        selected_letter = cands[0]["letter"]

    backup_ref = extract_candidate_key(backup_hyp)
    backup_letter = resolve_candidate_key(backup_ref, cands)

    falsify = (
        main_pkg.get("证伪条件")
        or main_pkg.get("证伪/退出条件")
        or main_pkg.get("证伪/削弱条件")
        or [""]
    )[0]
    uncertainty = (main_pkg.get("不确定性声明") or main_pkg.get("不确定性") or [""])[0]

    scores = parse_scores(header, total, secs.get(4, ""), warnings=warnings)
    unscored = [c["letter"] for c in cands if c["letter"] not in scores]
    if unscored:
        warnings.append("候选总分未结构化：" + "、".join(unscored) + "（保留为 —）")

    cand_name = ""
    for c in cands:
        if c["letter"] == selected_letter:
            cand_name = c["name"]
            break
    if cand_name:
        cand_name = re.sub(r"[「『\"（(].*$", "", cand_name).strip(" ·-—：:")
    paper_title = (sec10.get("paper_title") or [""])[
        0
    ].strip() or f"{qid} {cand_name or safe_clip(question, 50)}：科学假说与研究计划设计"

    if not main_statement and selected_letter:
        selected = next((c for c in cands if c["letter"] == selected_letter), None)
        if selected:
            main_statement = selected["statement"]

    abstract_parts = []
    if main_statement:
        abstract_parts.append(main_statement)
    elif main_hyp:
        abstract_parts.append(main_hyp)
    if competing_explanation:
        abstract_parts.append("竞争解释：" + safe_clip(competing_explanation, 240))
    if predictions:
        abstract_parts.append(
            "主要可检验预测：" + "；".join(safe_clip(p, 160) for p in predictions[:2])
        )
    if falsify:
        abstract_parts.append("证伪条件：" + safe_clip(falsify, 160))
    abstract_parts.append("本研究计划为待执行验证方案，未执行任何实验。")
    paper_abstract = safe_clip(" ".join(abstract_parts), 1200)

    dataset_targets = build_dataset_targets(plan_rows)
    if not (main_statement or main_hyp):
        warnings.append("rationale 为空（需人工补录主假说陈述）")
    if any(not s.get("source_url") for s in sources):
        warnings.append("部分来源缺少可转换的 URL/DOI（locator 已保留）")
    if (
        dataset_targets
        and dataset_targets[0].get("status") == "proposed_target_not_structured"
    ):
        warnings.append("datasets.target 未结构化（已写入明确缺失状态）")

    L = []
    L.append(f"# {qid} 逐题结果文档")
    L.append("")
    L.append(f"- 题目原文：{question}")
    L.append(f"- 领域：{domain}")
    L.append(
        "- 流程状态：已完成完整挑战杯假说生成流程（问题解析、证据建联、多候选生成、七维评审、反馈修订、主备筛选与研究计划）。"
    )
    L.append(
        "- 结果类型：generated_hypothesis + proposed_research_plan；研究计划为待执行验证方案。"
    )
    L.append("")
    L.append("## 1. 案例的科学内容（对应模板 P15/表15）")
    L.append("")
    L.append("| 内容 | 本题实际情况 |")
    L.append("|---|---|")
    L.append(f"| 科学问题原文 | {esc_cell(question)} |")
    L.append(
        f"| 研究对象与关键变量 | {esc_cell((s1f.get('复述') or s1f.get('问题复述') or [''])[0] or secs.get(1, ''), 500)} |"
    )
    supp = "；".join(
        f"{s['sid']} {safe_clip(s['title'], 100)}：{safe_clip(s['contribution'], 80)}"
        for s in sources[:3]
    )
    L.append(f"| 已有认识与主要证据 | 共 {len(sources)} 条来源。{esc_cell(supp)} |")
    gap = safe_clip(uncertainty or (cands[0]["risk"] if cands else ""), 300)
    L.append(f"| 尚未解决的知识缺口 | {esc_cell(gap) or '见第 6 节不确定性'} |")
    L.append(
        f"| 需要保留的不确定性或争议 | {esc_cell((s1f.get('题意存疑点') or s1f.get('题意存疑') or [''])[0] or uncertainty, 300) or '—'} |"
    )
    L.append("")

    L.append("## 2. 候选假设与第一轮处理结果（对应模板 P15/表16）")
    L.append("")
    L.append(f"- 最终主假说：{safe_clip(main_hyp, 500)}")
    L.append(f"- 最终备选假说：{safe_clip(backup_hyp, 500)}")
    L.append("")
    L.append(
        "| 候选假设 | 主要依据 | 反对证据或替代解释 | 可检验预测 | 第一轮处理结果 |"
    )
    L.append("|---|---|---|---|---|")
    for i, c in enumerate(cands, 1):
        preds = "；".join(c["predictions"]) if c["predictions"] else ""
        if not preds and c["letter"] == selected_letter:
            preds = (
                "；".join(safe_clip(p, 120) for p in predictions[:3])
                or "见主假说包预测"
            )
        if not preds:
            preds = "未单列（非选中候选，保留为审计对照）"
        if c["letter"] == selected_letter:
            selected_note = (
                f"（最终筛选：候选{selected_ref}）"
                if selected_ref and selected_ref != selected_letter
                else ""
            )
            outcome = (
                f"七维总分 {scores.get(c['letter'], '—')}；"
                f"选中为主假说{selected_note}"
            )
        elif backup_letter and c["letter"] == backup_letter:
            backup_note = (
                f"（最终筛选：候选{backup_ref}）"
                if backup_ref and backup_ref != backup_letter
                else ""
            )
            outcome = (
                f"七维总分 {scores.get(c['letter'], '—')}；"
                f"保留为备选假说{backup_note}"
            )
        elif c["letter"] not in scores:
            outcome = "七维总分 —；落选（正文无该候选逐维评分依据，保留为对照）"
        else:
            outcome = f"七维总分 {scores.get(c['letter'], '—')}；落选"
        L.append(
            f"| H-{i:02d}（候选{c['letter']}·{esc_cell(c['name'], 40)}） | {esc_cell(c['evidence'] or c['raw'], 260)} | {esc_cell(c['risk'], 260) or '—'} | {esc_cell(preds, 320)} | {esc_cell(outcome)} |"
        )
    L.append("")

    L.append("## 3. 研究计划（待执行验证方案，对应模板 P15/表17）")
    L.append("")
    L.append("| 关键步骤 | 实际计划内容 | 该步骤能够支持、反对或区分什么 |")
    L.append("|---|---|---|")
    for cells in plan_rows:
        if len(cells) >= 4:
            L.append(
                f"| {esc_cell(cells[0])} | {esc_cell(cells[1])}（数据/实验：{esc_cell(cells[2])}） | {esc_cell(cells[3])} |"
            )
        elif len(cells) == 3:
            L.append(
                f"| {esc_cell(cells[0])} | {esc_cell(cells[1])} | {esc_cell(cells[2])} |"
            )
    L.append("")

    L.append("## 4. 第一轮问题与第二轮调整（对应模板 P16/表19）")
    L.append("")
    L.append(
        "| 第一轮具体问题 | 判断依据 | 对科学结论或研究计划的影响 | 第二轮实际调整 |"
    )
    L.append("|---|---|---|---|")
    for rv in revisions:
        trig = rv.get("trigger") or "（见修订原文附录）"
        L.append(
            f"| {esc_cell(trig, 240) or '评审/证据触发'} | 证据或评审缺陷（触发来源见原文） | 修订前：{esc_cell(rv.get('before') or '—', 240)} | 修订后：{esc_cell(rv.get('after') or '—', 240)} |"
        )
    L.append("")

    L.append("## 5. 两轮变化对照（对应模板 P17/表21）")
    L.append("")
    L.append("| 变化内容 | 第一轮 | 第二轮 | 变化原因 | 实际结果 |")
    L.append("|---|---|---|---|---|")
    for rv in revisions:
        cat = revision_category(rv.get("trigger"))
        L.append(
            f"| {esc_cell(cat)} | {esc_cell(rv.get('before') or '—', 200)} | {esc_cell(rv.get('after') or '—', 200)} | {esc_cell(rv.get('trigger') or '—', 160)} | {esc_cell(rv.get('why') or '见修订记录', 200)} |"
        )
    L.append("")

    L.append("## 6. 任务书标准字段（CompetitionResultView 投影）")
    L.append("")
    L.append(
        f"- problem_statement：{safe_clip((s1f.get('复述') or s1f.get('问题复述') or [''])[0] or question, 500)}"
    )
    L.append(f"- rationale：{main_statement or main_hyp}")
    if competing_explanation:
        L.append(f"- competing_explanation：{safe_clip(competing_explanation, 500)}")
    L.append(
        f"- technical_details：研究计划共 {len(plan_rows)} 步（见第 3 节）；候选 {len(cands)} 个、实质修订 {len(revisions)} 次、来源 {len(sources)} 条。"
    )
    L.append(f"- paper_title：{paper_title}")
    L.append(f"- paper_abstract：{paper_abstract}")
    L.append(
        "- experiments.execution_mode：proposed（全部为待执行验证方案，未执行实验）"
    )
    L.append("- results.classification：expected（预期结果，非实际实验结果）")
    L.append(f"- datasets.source：{len(sources)} 条来源。")
    L.append("")
    L.append("| 证据 ID | 关系 | 核验状态 | 来源 | 定位 | 具体作用 |")
    L.append("|---|---|---|---|---|---|")
    for s in sources:
        relation = infer_source_relation(s["stype"], s["contribution"])
        verification = infer_verification_status(s.get("source_url"))
        locator = s.get("source_url") or s["url"] or "未提供可解析定位符"
        L.append(
            f"| {esc_cell(s['sid'])} | {relation} | {verification} | "
            f"{esc_cell(s['title'], 180)} | {esc_cell(locator, 220)} | "
            f"{esc_cell(s['contribution'], 320)} |"
        )
    L.append("")

    md_doc = "\n".join(L)

    proj = {
        "projection_type": "competition_result_view",
        "projection_version": "content-layer-v1",
        "catalog_id": "science-125-questions-2021",
        "question_id": qid,
        "question_en": question,
        "status": "hypothesis_generation_complete",
        "workflow_completed": True,
        "workflow_scope": [
            "problem_analysis",
            "evidence_grounding",
            "candidate_generation",
            "seven_dimension_review",
            "feedback_revision",
            "selection",
            "research_plan",
        ],
        "generated_at": now_iso(),
        "result_classification": {
            "generated_hypothesis": True,
            "proposed_research_plan": True,
            "expected_result": True,
            "actual_execution_performed": False,
        },
        "competition_result_view": {
            "problem_statement": safe_clip(
                (s1f.get("复述") or s1f.get("问题复述") or [""])[0] or question, 2000
            ),
            "rationale": main_statement or main_hyp,
            "competing_explanation": competing_explanation or None,
            "technical_details": f"研究计划 {len(plan_rows)} 步；候选 {len(cands)} 个；修订 {len(revisions)} 次；来源 {len(sources)} 条。",
            "paper_title": paper_title,
            "paper_abstract": paper_abstract,
            "datasets": {
                "source": [
                    {
                        "evidence_id": s["sid"],
                        "title": s["title"],
                        "source_url": s.get("source_url"),
                        "publication_note": s["year"],
                        "relation": infer_source_relation(
                            s["stype"], s["contribution"]
                        ),
                        "verification_status": infer_verification_status(
                            s.get("source_url")
                        ),
                        "contribution": s["contribution"],
                    }
                    for s in sources
                ],
                "target": dataset_targets,
            },
            "methods": {
                "candidate_generation": "complete challenge-cup hypothesis generation workflow",
                "review": "seven-dimension review",
                "selection": {
                    "primary": main_hyp,
                    "backup": backup_hyp,
                    "primary_ref": selected_ref or None,
                    "backup_ref": backup_ref or None,
                },
                "workflow_completed": True,
            },
            "experiments": {
                "execution_mode": "proposed",
                "proposed_steps": len(plan_rows),
            },
            "results": {
                "classification": "expected",
                "status": "proposed",
                "decision": None,
                "summary": safe_clip(paper_abstract, 600),
                "artifact_refs": [f"summaries/{qid}.md"],
            },
            "references": [
                {
                    "ref_id": s["sid"],
                    "title": s["title"],
                    "locator": s["url"],
                    "authors": s["authors"],
                    "year": s["year"],
                }
                for s in sources
            ],
        },
        "conversion_warnings": warnings,
    }

    stats = {
        "question": question,
        "workflow_completed": True,
        "sources": len(sources),
        "cands": len(cands),
        "scored_cands": len([c for c in cands if c["letter"] in scores]),
        "revisions": len(substantive_revisions),
        "plan_steps": len(plan_rows),
        "predictions": len(predictions),
        "selection_complete": bool(main_hyp and backup_hyp),
        "falsification_complete": bool(falsify),
    }
    return md_doc, proj, warnings, stats


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--challenge-root", type=Path, default=DEFAULT_CHALLENGE_ROOT)
    parser.add_argument("--source-dir", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--ledger", type=Path)
    parser.add_argument("--catalog", type=Path)
    return parser.parse_args(argv)


def _assert_fresh_output_dir(output_dir):
    if output_dir.exists() and any(output_dir.iterdir()):
        raise RuntimeError(f"output directory must be empty: {output_dir}")
    output_dir.mkdir(parents=True, exist_ok=True)


def _assert_no_private_metadata(text, *, label):
    forbidden = (
        r"[A-Za-z]:\\Users\\",
        r"执行\s*Agent",
        r"编排方\s*agent",
        r"\b(?:instance|session)[_-]?id\b",
    )
    for pattern in forbidden:
        if re.search(pattern, text, re.IGNORECASE):
            raise RuntimeError(f"private metadata detected in {label}: {pattern}")


def validate_export_package(output_dir):
    output_dir = Path(output_dir).resolve()
    manifest_path = output_dir / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    documents = manifest.get("documents") or []
    expected_ids = [f"SCI-{index:03d}" for index in range(1, 126)]
    observed_ids = [item.get("question_id") for item in documents]
    if observed_ids != expected_ids:
        raise RuntimeError("manifest must contain SCI-001 through SCI-125 in order")
    allowed_relations = {"supports", "challenges", "context", "method", "boundary"}
    allowed_verification = {
        "unverified",
        "metadata_checked",
        "full_text_checked",
        "human_verified",
    }
    hash_checks = 0
    source_count = 0
    for item in documents:
        qid = item["question_id"]
        if item.get("status") != "hypothesis_generation_complete":
            raise RuntimeError(f"{qid}: workflow status is incomplete")
        if item.get("candidates_count", 0) < 2:
            raise RuntimeError(f"{qid}: fewer than two candidates")
        if item.get("scored_candidates_count", 0) < item.get("candidates_count", 0):
            raise RuntimeError(
                f"{qid}: not every candidate has a structured review score"
            )
        if item.get("revisions_count", 0) < 1:
            raise RuntimeError(f"{qid}: missing substantive revision")
        if item.get("plan_steps", 0) < 1:
            raise RuntimeError(f"{qid}: missing research plan")
        if item.get("predictions_count", 0) < 1:
            raise RuntimeError(f"{qid}: missing testable prediction")
        if item.get("selection_complete") is not True:
            raise RuntimeError(f"{qid}: primary/backup selection is incomplete")
        if item.get("falsification_complete") is not True:
            raise RuntimeError(f"{qid}: falsification condition is missing")
        summary_path = output_dir / item["summary_path"]
        projection_path = output_dir / item["projection_path"]
        if sha256_file(summary_path) != item["summary_sha256"]:
            raise RuntimeError(f"{qid}: summary hash mismatch")
        if sha256_file(projection_path) != item["projection_sha256"]:
            raise RuntimeError(f"{qid}: projection hash mismatch")
        hash_checks += 2
        summary = summary_path.read_text(encoding="utf-8")
        projection_text = projection_path.read_text(encoding="utf-8")
        _assert_no_private_metadata(summary, label=item["summary_path"])
        _assert_no_private_metadata(projection_text, label=item["projection_path"])
        if "附录：原始假说包全文" in summary:
            raise RuntimeError(f"{qid}: raw source appendix must not be exported")
        for section in range(1, 7):
            if f"## {section}." not in summary:
                raise RuntimeError(f"{qid}: summary section {section} is missing")
        projection = json.loads(projection_text)
        if projection.get("status") != "hypothesis_generation_complete":
            raise RuntimeError(f"{qid}: projection workflow status is incomplete")
        classification = projection.get("result_classification") or {}
        actual = classification.get("actual_execution_performed") is True
        expected = classification.get("expected_result") is True
        if actual == expected:
            raise RuntimeError(
                f"{qid}: actual and expected result flags must be exclusive"
            )
        sources = (
            projection.get("competition_result_view", {})
            .get("datasets", {})
            .get("source", [])
        )
        if len(sources) < 4:
            raise RuntimeError(f"{qid}: fewer than four evidence sources")
        source_count += len(sources)
        for source in sources:
            if source.get("relation") not in allowed_relations:
                raise RuntimeError(f"{qid}: unsupported evidence relation")
            if source.get("verification_status") not in allowed_verification:
                raise RuntimeError(f"{qid}: unsupported verification status")
            source_url = source.get("source_url")
            if source_url and not re.match(r"^https?://", source_url, re.IGNORECASE):
                raise RuntimeError(f"{qid}: malformed source URL")
    return {
        "status": "passed",
        "documents": len(documents),
        "hash_checks": hash_checks,
        "evidence_sources": source_count,
    }


def export_results(*, source_dir, output_dir, ledger_path, catalog_path):
    source_dir = Path(source_dir).resolve()
    output_dir = Path(output_dir).resolve()
    ledger_path = Path(ledger_path).resolve()
    catalog_path = Path(catalog_path).resolve()
    _assert_fresh_output_dir(output_dir)
    sum_dir = output_dir / "summaries"
    proj_dir = output_dir / "projections"
    sum_dir.mkdir()
    proj_dir.mkdir()
    catalog = load_catalog_questions(catalog_path)
    done = ledger_done_questions(ledger_path)
    converted = 0
    all_warnings = []
    docs = []

    files = sorted(
        f.name for f in source_dir.iterdir() if re.match(r"^SCI-\d{3}\.md$", f.name)
    )
    for f in files:
        qid = f[:-3]
        src_path = source_dir / f
        if qid not in done or src_path.stat().st_size == 0:
            continue
        src_text = src_path.read_text(encoding="utf-8")
        try:
            md_doc, proj, warnings, stats = build_outputs(
                qid, src_path, catalog.get(qid), src_text
            )
        except (
            IndexError,
            KeyError,
            OSError,
            RuntimeError,
            TypeError,
            ValueError,
        ) as e:
            all_warnings.append(f"{qid}: 转换异常 {e}")
            continue
        for w in warnings:
            all_warnings.append(f"{qid}: {w}")
        projection_text = json.dumps(proj, ensure_ascii=False, indent=2)
        _assert_no_private_metadata(md_doc, label=f"summaries/{qid}.md")
        _assert_no_private_metadata(projection_text, label=f"projections/{qid}.json")
        sum_path = sum_dir / (qid + ".md")
        proj_path = proj_dir / (qid + ".json")
        sum_path.write_text(md_doc, encoding="utf-8")
        proj_path.write_text(projection_text, encoding="utf-8")
        converted += 1
        docs.append(
            {
                "question_id": qid,
                "question_en": stats["question"],
                "status": "hypothesis_generation_complete",
                "workflow_completed": stats["workflow_completed"],
                "summary_path": f"summaries/{qid}.md",
                "summary_sha256": sha256_file(sum_path),
                "projection_path": f"projections/{qid}.json",
                "projection_sha256": sha256_file(proj_path),
                "sources_count": stats["sources"],
                "candidates_count": stats["cands"],
                "scored_candidates_count": stats["scored_cands"],
                "revisions_count": stats["revisions"],
                "plan_steps": stats["plan_steps"],
                "predictions_count": stats["predictions"],
                "selection_complete": stats["selection_complete"],
                "falsification_complete": stats["falsification_complete"],
                "warnings": warnings,
            }
        )

    docs.sort(key=lambda d: d["question_id"])
    manifest = {
        "manifest_version": "2",
        "catalog_id": "science-125-questions-2021",
        "generated_at": now_iso(),
        "layer": "hypothesis_generation",
        "status": ("partial" if converted < 125 else "hypothesis_generation_complete"),
        "declaration": (
            f"{converted}/125 题均已完成完整挑战杯假说生成流程。"
            if converted == 125
            else f"已完成 {converted}/125 题的挑战杯假说生成流程。"
        ),
        "counts": {
            "total_questions": 125,
            "documents": converted,
            "workflow_completed": converted,
            "not_generated": 125 - converted,
            "documents_with_warnings": sum(1 for d in docs if d["warnings"]),
            "warnings": len(all_warnings),
        },
        "warnings": all_warnings,
        "documents": docs,
    }
    if converted != 125:
        raise RuntimeError(f"incomplete catalog export: {converted}/125")
    man_path = output_dir / "manifest.json"
    manifest_text = json.dumps(manifest, ensure_ascii=False, indent=2)
    _assert_no_private_metadata(manifest_text, label="manifest.json")
    man_path.write_text(manifest_text, encoding="utf-8")
    manifest["validation"] = validate_export_package(output_dir)
    man_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print(f"转换完成：{converted}/125 题")
    print(f"manifest: {man_path}")
    if all_warnings:
        print(f"警告 {len(all_warnings)} 条：")
        for w in all_warnings:
            print("  -", w)
    else:
        print("无解析警告")
    return manifest


def main(argv=None):
    args = parse_args(argv)
    challenge_root = args.challenge_root.resolve()
    export_results(
        source_dir=args.source_dir or challenge_root / "08-高质量假说结果",
        output_dir=args.output_dir,
        ledger_path=args.ledger or challenge_root / "125题高质量假说进度账本.xlsx",
        catalog_path=args.catalog
        or PROJECT_ROOT
        / "core"
        / "research"
        / "competition"
        / "data"
        / "science_125_questions.json",
    )


if __name__ == "__main__":
    main()
